"""
Pulls the quarterly affordable-lettings series out of the Homes Victoria workbook
into app/data/source/lettings_series.json, which scripts/build-data.mjs then merges.

The "lga aff total" sheet's Percent column for Sep 2025 matches affordability_pct in
master_data_v2.json exactly for all 79 areas, so this is the same source the data team used.

Run:
  data_pipeline/.venv/bin/python3 scripts/extract-lettings-series.py \
    "~/Downloads/Affordable rental dwellings by Local Government Area - September quarter 2025.xlsx"
"""
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIR = os.path.join(HERE, "..", "app", "data", "source")
OUT = os.path.join(SOURCE_DIR, "lettings_series.json")

SHEET = "lga aff total"
QUARTERS = 21  # Sep 2020 to Sep 2025 inclusive, the "last 5 years" the copy promises
NOT_AN_LGA = {"Metro", "Non-Metro", "Table Total"}

# The workbook's spellings against the names in master_data_v2.json.
NAME_FIXES = {
    "Colac-Otway": "Colac Otway",
    "Merri-bek": "Moreland",
    "Mornington Penin'a": "Mornington Peninsula",
}


def fail(message):
    sys.exit(f"extract-lettings-series: {message}")


def main():
    if len(sys.argv) < 2:
        fail("pass the path to the Homes Victoria xlsx")
    path = os.path.expanduser(sys.argv[1])
    if not os.path.exists(path):
        fail(f"no such file: {path}")

    with open(os.path.join(SOURCE_DIR, "master_data_v2.json")) as f:
        codes = {row["lga_name"]: row["lga_code"] for row in json.load(f)}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        fail(f'"{SHEET}" is missing; the workbook has {wb.sheetnames}')
    rows = list(wb[SHEET].iter_rows(values_only=True))

    # Row 3 names each quarter across a merged pair of columns, row 4 labels the pair.
    header, sub = rows[2], rows[3]
    quarter_of, current = {}, None
    for i in range(1, len(header)):
        if header[i]:
            current = str(header[i]).strip()
        quarter_of[i] = current
    percent_cols = [(i, quarter_of[i]) for i in range(1, len(sub)) if sub[i] == "Percent"]
    if not percent_cols:
        fail("found no Percent columns")
    percent_cols = percent_cols[-QUARTERS:]
    quarters = [q for _, q in percent_cols]

    series, unmatched = {}, []
    for row in rows[4:]:
        name = (row[0] or "").strip()
        if not name or name in NOT_AN_LGA:
            continue
        name = NAME_FIXES.get(name, name)
        code = codes.get(name)
        if code is None:
            unmatched.append(name)
            continue
        values = [row[i] for i, _ in percent_cols]
        if any(v is None for v in values):
            fail(f"{name} has a gap in the last {QUARTERS} quarters")
        # The sheet stores proportions; the app shows percentages.
        series[str(code)] = [round(float(v) * 100, 1) for v in values]

    if unmatched:
        fail(f"these sheet rows match no area in master_data_v2.json: {unmatched}")
    if len(series) != len(codes):
        missing = [n for n, c in codes.items() if str(c) not in series]
        fail(f"expected {len(codes)} areas, got {len(series)}; missing {missing}")

    out = {
        "_source": (
            "Homes Victoria, Affordable rental dwellings by Local Government Area, "
            "September quarter 2025. Sheet 'lga aff total', Percent columns."
        ),
        "quarters": quarters,
        "series": series,
    }
    with open(OUT, "w") as f:
        json.dump(out, f, indent=1)
        f.write("\n")
    print(
        f"extract-lettings-series: wrote {len(series)} areas x {len(quarters)} quarters "
        f"({quarters[0]} to {quarters[-1]}) to app/data/source/lettings_series.json"
    )


if __name__ == "__main__":
    main()
