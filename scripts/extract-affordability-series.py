"""
Pulls the quarterly affordable-lettings series out of the Homes Victoria workbook
in data_pipeline/iteration_2_data into app/data/affordability-series.json.

The workbook is the source the team's CSVs were cut down from: for all 79 areas the
"lga aff total" sheet's Sep 2025 Percent equals affordability_pct in master_data_v4.json
exactly, and Sep 2025 minus Sep 2022 equals affordability_trend exactly. So the chart
draws the history of the same number the ranking scores.

Five series per area (all bedrooms, then 1/2/3/4 bedrooms), 103 quarters each,
Mar 2000 to Sep 2025. Both the affordable count and the share are kept: the count is
what makes a pooled multi-quarter share possible for the small areas whose
quarter-to-quarter share swings wildly.

Shares are stored as integers in tenths of a percent (0.798 -> 798). The workbook
carries at most three decimal places on the fraction, so this is lossless and keeps
the file about a third smaller than decimal strings would.

Run:
  data_pipeline/.venv/bin/python scripts/extract-affordability-series.py
"""
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
WORKBOOK = os.path.join(
    ROOT,
    "data_pipeline",
    "iteration_2_data",
    "Affordable_rental_dwellings_by_LGA_SEP-Quarter-2025.xlsx",
)
MASTER = os.path.join(ROOT, "data_pipeline", "iteration_2_data", "master_data_v4.json")
OUT = os.path.join(ROOT, "app", "data", "affordability-series.json")

# Sheet name -> the key used in the output.
SHEETS = {
    "lga aff total": "all",
    "lga aff 1br": "br1",
    "lga aff 2br": "br2",
    "lga aff 3br": "br3",
    "lga aff 4br": "br4",
}

EXPECTED_AREAS = 79
EXPECTED_QUARTERS = 103
FIRST_QUARTER = "Mar 2000"
LAST_QUARTER = "Sep 2025"

HEADER_ROW = 3        # quarter name, merged across the Affordable/Percent pair
SUBHEADER_ROW = 4     # "Affordable" / "Percent"
FIRST_DATA_ROW = 5

# Rows that are totals rather than areas.
NOT_AN_AREA = {"Metro", "Non-Metro", "Table Total"}

# The workbook's spellings, against the names the data team uses.
WORKBOOK_NAMES = {
    "Colac-Otway": "Colac Otway",
    "Mornington Penin'a": "Mornington Peninsula",
}
# The council renamed itself in 2022; the workbook uses the new name, the team's data the old one.
MASTER_NAMES = {"Moreland": "Merri-bek"}


def fail(message):
    sys.exit(f"extract-affordability-series: {message}")


def quarter_columns(ws, sheet_name):
    """The (count_col, percent_col, quarter_label) triples, left to right."""
    label_of, current = {}, None
    for col in range(2, ws.max_column + 1):
        value = ws.cell(HEADER_ROW, col).value
        if value not in (None, "", " "):
            current = str(value).strip()
        label_of[col] = current

    triples, pending = [], {}
    for col in range(2, ws.max_column + 1):
        kind = ws.cell(SUBHEADER_ROW, col).value
        label = label_of[col]
        if kind == "Affordable":
            pending[label] = col
        elif kind == "Percent":
            if label not in pending:
                fail(f"{sheet_name}: a Percent column for {label} has no Affordable column")
            triples.append((pending.pop(label), col, label))
    if pending:
        fail(f"{sheet_name}: Affordable columns with no Percent: {sorted(pending)}")
    if len(triples) != EXPECTED_QUARTERS:
        fail(f"{sheet_name}: expected {EXPECTED_QUARTERS} quarters, found {len(triples)}")
    return triples


def tenths_of_percent(fraction, where):
    """0.798 -> 798. Rejects anything the workbook shouldn't contain."""
    if not isinstance(fraction, (int, float)):
        fail(f"{where}: share is {fraction!r}, expected a number")
    if not 0 <= fraction <= 1:
        fail(f"{where}: share {fraction} is outside 0-1")
    scaled = fraction * 1000
    nearest = round(scaled)
    if abs(scaled - nearest) > 1e-6:
        fail(f"{where}: share {fraction} needs more precision than 0.1 of a percent")
    return int(nearest)


def read_sheet(wb, sheet_name, key, codes):
    if sheet_name not in wb.sheetnames:
        fail(f'sheet "{sheet_name}" is missing; the workbook has {wb.sheetnames}')
    ws = wb[sheet_name]
    triples = quarter_columns(ws, sheet_name)
    labels = [label for _, _, label in triples]
    if labels[0] != FIRST_QUARTER or labels[-1] != LAST_QUARTER:
        fail(f"{sheet_name}: quarters run {labels[0]} to {labels[-1]}, expected {FIRST_QUARTER} to {LAST_QUARTER}")

    out, seen = {}, set()
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        raw = ws.cell(row, 1).value
        name = str(raw).strip() if raw is not None else ""
        if not name or name in NOT_AN_AREA:
            continue
        name = WORKBOOK_NAMES.get(name, name)
        code = codes.get(name)
        if code is None:
            fail(f'{sheet_name} row {row}: "{name}" matches no area in master_data_v4.json')
        if code in seen:
            fail(f"{sheet_name}: {name} appears twice")
        seen.add(code)

        counts, shares = [], []
        for count_col, percent_col, label in triples:
            where = f"{sheet_name} {name} {label}"
            count = ws.cell(row, count_col).value
            if not isinstance(count, (int, float)) or count < 0 or count != int(count):
                fail(f"{where}: affordable count is {count!r}, expected a whole number")
            counts.append(int(count))
            shares.append(tenths_of_percent(ws.cell(row, percent_col).value, where))
        out[code] = {"n": counts, "pct": shares}

    if len(out) != EXPECTED_AREAS:
        fail(f"{sheet_name}: matched {len(out)} areas, expected {EXPECTED_AREAS}")
    return labels, out


def main():
    if not os.path.exists(WORKBOOK):
        fail(f"no workbook at {WORKBOOK}")
    with open(MASTER) as f:
        master = json.load(f)
    codes = {MASTER_NAMES.get(r["lga_name"], r["lga_name"]): r["lga_code"] for r in master}
    if len(codes) != EXPECTED_AREAS:
        fail(f"master_data_v4.json has {len(codes)} unique names, expected {EXPECTED_AREAS}")

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    quarters = None
    by_key = {}
    for sheet_name, key in SHEETS.items():
        labels, series = read_sheet(wb, sheet_name, key, codes)
        if quarters is None:
            quarters = labels
        elif labels != quarters:
            fail(f"{sheet_name}: its quarters differ from {list(SHEETS)[0]}")
        by_key[key] = series

    # Pivot to area -> bedroom key -> {n, pct}, so a page loads one area's block.
    series = {}
    for code in sorted(by_key["all"]):
        series[str(code)] = {key: by_key[key][code] for key in SHEETS.values()}

    # The scored number must be the last point of the all-bedrooms share.
    last = quarters.index(LAST_QUARTER)
    for row in master:
        code = str(row["lga_code"])
        got = series[code]["all"]["pct"][last] / 1000
        if abs(got - row["affordability_pct"]) > 5e-4:
            fail(
                f"{row['lga_name']}: workbook {LAST_QUARTER} share {got} does not match "
                f"affordability_pct {row['affordability_pct']} in master_data_v4.json"
            )

    out = {
        "_source": (
            "Homes Victoria, Affordable rental dwellings by Local Government Area, "
            "September quarter 2025. One series per bedroom size; 'all' is the workbook's "
            "total sheet. 'n' is the count of new lettings that were affordable; 'pct' is "
            "that share in tenths of a percent."
        ),
        "quarters": quarters,
        "bedrooms": list(SHEETS.values()),
        "series": series,
    }
    with open(OUT, "w") as f:
        json.dump(out, f, separators=(",", ":"), sort_keys=False)
        f.write("\n")

    size = os.path.getsize(OUT)
    print(
        f"extract-affordability-series: wrote {len(series)} areas x {len(SHEETS)} series "
        f"x {len(quarters)} quarters ({quarters[0]} to {quarters[-1]}) "
        f"to app/data/affordability-series.json ({size / 1024:.0f} KB)"
    )


if __name__ == "__main__":
    main()
